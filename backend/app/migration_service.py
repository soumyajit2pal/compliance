"""Immutable source snapshots, profiling, content validation and Draft-only loads."""
import csv,io,json,hashlib,zipfile
from datetime import date,datetime
from collections import Counter
import sqlalchemy as sa
from fastapi import HTTPException
from openpyxl import load_workbook
from . import models as m
from .events import new_id,now,audit
from . import repository as repo
from .schemas import InstrumentInput,ClauseInput
from .access import actor_for
from .import_template import canonical_header,LABELS,RECORDS_SHEET,GUIDANCE_SHEETS,GUIDE_TITLE
from .import_choices import current_master,materiality_code

MAX_ROWS=10000

def parse_source(name,data):
    rows=[]
    if name.lower().endswith('.csv'):
        try:
            reader=csv.DictReader(io.StringIO(data.decode('utf-8-sig')))
            if not reader.fieldnames:raise ValueError('Missing headers')
            headers=[canonical_header(h) for h in reader.fieldnames]
            if len(headers)>100 or any(not h for h in headers) or len(set(headers))!=len(headers):raise ValueError('Invalid or duplicate headers')
            for number,row in enumerate(reader,2):
                if None in row:raise ValueError('Extra cells without a header')
                rows.append(('CSV',number,row))
                if len(rows)>MAX_ROWS: raise ValueError('Too many rows')
        except (UnicodeError,csv.Error,ValueError) as exc: raise HTTPException(422,'Use a UTF-8 CSV with column headers.') from exc
    elif name.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist())>50*1024*1024:raise ValueError('Workbook too large')
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            try:
                template_guide = book['How to fill'] if 'How to fill' in book.sheetnames else None
                is_template = template_guide is not None and template_guide['A1'].value == GUIDE_TITLE
                if is_template and RECORDS_SHEET not in book.sheetnames:
                    raise ValueError('The Regulatory records sheet is missing')
                if is_template and set(book.sheetnames) - ({RECORDS_SHEET} | GUIDANCE_SHEETS):
                    raise ValueError('Move all records into the Regulatory records sheet')
                for sheet in book:
                    if is_template and sheet.title != RECORDS_SHEET:continue
                    iterator=sheet.iter_rows(values_only=True);headers=next(iterator,())
                    if not headers:continue
                    if len(headers)>100:raise ValueError('Too many columns')
                    canonical=[canonical_header(h) for h in headers]
                    if any(not h for h in canonical) or len(set(canonical))!=len(canonical):raise ValueError('Invalid or duplicate headers')
                    for number,values in enumerate(iterator,2):
                        if any(v is not None for v in values): rows.append((sheet.title,number,dict(zip([str(h or '') for h in headers],values))))
                        if len(rows)>MAX_ROWS:raise ValueError('Too many rows')
            finally:book.close()
        except Exception as exc:raise HTTPException(422,'The XLSX could not be read safely. Keep the template sheet names and put all records in Regulatory records. Limit the file to 10,000 rows and 50 MB expanded.') from exc
    else: raise HTTPException(422,'Choose a .csv or .xlsx source register.')
    if not rows or len(rows)>MAX_ROWS:raise HTTPException(422,'Enter between 1 and 10,000 records. In the Excel template, fill Regulatory records; the example and guidance sheets are not imported.')
    return rows


def normalise(db,raw,seen=None):
    """Raw payload never changes; canonical codes and findings are separate."""
    data={canonical_header(k):str(v).strip() if v is not None else '' for k,v in raw.items()}
    errors=[];normal={};kind=data.get('entity_type','instrument').lower()
    if len(data) != len(raw):errors.append('Duplicate columns refer to the same field; keep one column per screen field')
    def required(key):
        value=data.get(key,'')
        if not value: errors.append('Missing '+LABELS.get(key,key))
        return value
    def lookup(table,key,column='code',value_override=None):
        value=required(key) if value_override is None else str(value_override).strip()
        if table is m.organisational_unit and value.casefold()=='banking operations': value='OPERATIONS'
        query=sa.select(table).where(sa.func.upper(getattr(table.c,column))==value.upper())
        if table is m.user:
            display = table.c.full_name + sa.literal(' (') + table.c.username + sa.literal(')')
            query=sa.select(table).where(sa.or_(sa.func.upper(table.c.username)==value.upper(),sa.func.upper(display)==value.upper()))
        matches=db.execute(query).mappings().all() if value else []
        if not matches and value:
            name_column=table.c.full_name if table is m.user else table.c.name
            matches=db.execute(sa.select(table).where(sa.func.upper(name_column)==value.upper())).mappings().all()
        matches=[row for row in matches if table is m.regulatory_domain or current_master(row)]
        if len(matches)>1:
            errors.append('Ambiguous '+LABELS.get(key,key)+'; use the username or master code to identify one value');return None
        row=matches[0] if matches else None
        if not row:
            errors.append('Invalid '+LABELS.get(key,key));return None
        return row['id']
    def date_value(key,optional=False):
        value=data.get(key,'') if optional else required(key)
        if optional and not value:return None
        try:return datetime.fromisoformat(value).date().isoformat()
        except ValueError:pass
        for pattern in ('%d/%m/%Y','%d-%m-%Y'):
            try:return datetime.strptime(value,pattern).date().isoformat()
            except ValueError:pass
        errors.append('Invalid '+LABELS.get(key,key)+'; use DD/MM/YYYY or YYYY-MM-DD');return None
    if kind=='instrument':
        primary_domain_id=lookup(m.regulatory_domain,'domain_code')
        secondary_domain_ids=[]
        additional_domains=data.get('additional_domain_codes','')
        if additional_domains:
            tokens=[token.strip() for token in additional_domains.split(';')]
            if any(not token for token in tokens):
                errors.append('Additional affected domains must be non-empty values separated by semicolons')
            seen_domains=set()
            for token in (token for token in tokens if token):
                domain_id=lookup(m.regulatory_domain,'additional_domain_codes',value_override=token)
                if not domain_id: continue
                if domain_id==primary_domain_id:
                    errors.append('Additional affected domains must not repeat the primary domain')
                elif domain_id in seen_domains:
                    errors.append('Additional affected domains must not contain duplicates')
                else:
                    seen_domains.add(domain_id);secondary_domain_ids.append(domain_id)
        normal={'title':required('title'),'short_title':data.get('short_title') or None,
            'applicability_date':date_value('applicability_date',optional=True),'regulator_id':lookup(m.regulator,'regulator_code'),
            'instrument_type_id':lookup(m.instrument_type,'instrument_type_code'),
            'primary_domain_id':primary_domain_id,'secondary_domain_ids':secondary_domain_ids,
            'regulatory_owner_unit_id':lookup(m.organisational_unit,'department_code'),
            'interpretation_owner_user_id':lookup(m.user,'owner_username','username'),
            'confidentiality_class_id':lookup(m.confidentiality_class,'confidentiality_code'),
            'source_document_id':required('source_document_id'),'official_reference_number':data.get('official_reference_number') or None,
            'issue_date':date_value('issue_date'),'effective_date':date_value('effective_date'),
            'materiality':materiality_code(db,data.get('materiality','MEDIUM')),'keywords':data.get('keywords') or None}
        if not normal['materiality']:errors.append('Invalid Materiality; choose a current database value')
        if normal['source_document_id']:
            document=db.execute(sa.select(m.document).where(m.document.c.id==normal['source_document_id'])).mappings().first()
            if not document:errors.append('Missing source document')
            elif document['owner_unit_id']!=normal['regulatory_owner_unit_id'] or document['confidentiality_class_id']!=normal['confidentiality_class_id']:
                errors.append('Source document department/classification mismatch')
        if normal['interpretation_owner_user_id'] and normal['regulatory_owner_unit_id']:
            owner=actor_for(db,normal['interpretation_owner_user_id'])
            if not owner.covers(normal['regulatory_owner_unit_id'],{'AUTHOR'}):errors.append('Owner lacks scoped Author role')
            if owner.user['needs_email_confirmation_flag'] or owner.user['needs_department_selection_flag']:
                errors.append('Interpretation owner must complete the work profile')
        identity=(normal['regulator_id'],(normal['official_reference_number'] or normal['title']).casefold())
        if seen is not None:
            if identity in seen:errors.append('Duplicate instrument in source population')
            seen.add(identity)
        if normal['official_reference_number'] and db.execute(sa.select(m.regulatory_instrument.c.id).where(
            m.regulatory_instrument.c.regulator_id==normal['regulator_id'],sa.func.upper(sa.func.trim(m.regulatory_instrument.c.official_reference_number))==normal['official_reference_number'].upper())).first():errors.append('Official reference already registered')
        try:InstrumentInput(**normal)
        except Exception:errors.append('Invalid field length, materiality or required value')
        if any(word in data.get('title','').lower() for word in ['multiple obligations','combined obligations']):errors.append('Combined obligations require decomposition in Week 4')
        if data.get('title','').casefold().startswith('example only'):
            errors.append('Replace the example with an actual regulatory record before importing')
    elif kind=='clause':
        normal={'version_id':required('instrument_version_id'),'code':data.get('clause_reference') or None,
            'heading':data.get('heading') or None,'clause_text':required('clause_text'),
            'internal_interpretation':data.get('internal_interpretation') or None,
            'source_page_reference':data.get('source_page_reference') or None,
            'source_clause_reference':data.get('source_clause_reference') or None,'parent_clause_id':data.get('parent_clause_id') or None}
        version=db.execute(sa.select(m.instrument_version).where(m.instrument_version.c.id==normal['version_id'])).mappings().first()
        if not version or version['approval_status']!='DRAFT':errors.append('Target version must be an existing Draft')
        if normal['code'] and db.execute(sa.select(m.regulatory_clause.c.id).where(m.regulatory_clause.c.instrument_version_id==normal['version_id'],m.regulatory_clause.c.code==normal['code'])).first(): errors.append('Clause / Para reference already exists in target version')
        identity=(normal['version_id'],normal['code'])
        if seen is not None and normal['code']:
            if identity in seen:errors.append('Duplicate clause / para reference in source population')
            seen.add(identity)
    elif kind=='obligation':
        statement=required('statement');required('title');required('instrument_reference');required('clause_reference')
        unit=lookup(m.organisational_unit,'department_code');owner=lookup(m.user,'owner_username','username')
        if unit and owner and not actor_for(db,owner).covers(unit,{'BUSINESS_OWNER'}):errors.append('Invalid Business Owner for the mapped department')
        if ';' in statement or '\n' in statement or data.get('combined_obligation_flag','').lower() in {'1','true','yes'}:
            errors.append('Potential combined obligations: content owner must decompose and validate separate requirements')
        if data.get('process_name'):
            process=db.execute(sa.select(m.process.c.id).where(sa.or_(sa.func.upper(m.process.c.code)==data['process_name'].upper(),sa.func.upper(m.process.c.name)==data['process_name'].upper()),m.process.c.active_flag==1)).first()
            if not process:errors.append('Invalid process name; standardise against the process master')
        identity=(data.get('instrument_reference','').upper(),data.get('clause_reference','').upper(),statement.casefold())
        if seen is not None:
            if identity in seen:errors.append('Duplicate obligation statement and regulatory source')
            seen.add(identity)
    else:errors.append('Unrecognised entity_type; choose instrument, clause or obligation')
    if kind=='instrument':
        signature=[kind,normal.get('regulator_id'),' '.join((normal.get('official_reference_number') or normal.get('title') or '').split()).casefold()]
    elif kind=='clause':
        signature=[kind,normal.get('version_id'),normal.get('code') or [normal.get('source_page_reference'),normal.get('clause_text')]]
    else:
        signature=[kind,data.get('instrument_reference','').upper(),data.get('clause_reference','').upper(),' '.join(data.get('statement','').split()).casefold()]
    fingerprint=hashlib.sha256(json.dumps(signature,sort_keys=True).encode()).hexdigest()
    if seen is not None:
        if fingerprint in seen:errors.append('Duplicate source record in staged population')
        seen.add(fingerprint)
    return {'entity_type':kind,'normalised':normal,'errors':sorted(set(errors)),'fingerprint':fingerprint,
        'load_available':kind in {'instrument','clause'},'deferred_note':'Obligation loading and workflow are Week 4' if kind=='obligation' else None}


def stage(db,actor,filename,data,source_system,storage):
    actor.require({'AUTHOR'})
    if not any(g['role_code']=='AUTHOR' and g['global_scope_flag'] for g in actor.grants):raise HTTPException(403,'Migration staging requires a global Regulatory Author.')
    file_hash=hashlib.sha256(data).hexdigest()
    if db.execute(sa.select(m.migration_batch.c.id).where(m.migration_batch.c.file_hash==file_hash)).first():raise HTTPException(409,'This exact source file has already been staged. Open its existing batch.')
    rows=parse_source(filename,data);identifier=new_id();seen=set();at=now()
    for previous in db.execute(sa.select(m.migration_row.c.validation_errors).where(m.migration_row.c.validation_status!='LOADED')).scalars():
        profile=json.loads(previous or '{}')
        if profile.get('fingerprint'):seen.add(profile['fingerprint'])
    db.execute(m.migration_batch.insert().values(id=identifier,code='MIG-'+identifier[:12],source_file=filename,
        file_hash=file_hash,uploaded_by=actor.id,uploaded_at=at,source_count=len(rows),pending_count=len(rows),record_status='PROFILED'))
    for sheet,number,raw in rows:
        result=normalise(db,raw,seen)
        # A physical row identity retains every occurrence, including duplicates.
        legacy=hashlib.sha256((file_hash+':'+sheet+':'+str(number)).encode()).hexdigest()
        db.execute(m.migration_row.insert().values(id=new_id(),batch_id=identifier,sheet_name=sheet,row_number=number,
            source_system=source_system,legacy_id=legacy,source_payload=json.dumps(raw,default=str),
            validation_status='CORRECTION_REQUIRED' if result['errors'] else 'CONTENT_VALIDATION_PENDING',validation_errors=json.dumps(result)))
    from .storage import stored_path, write_private
    path=stored_path(storage,identifier,staging=True)
    path.parent.mkdir(mode=0o700,parents=True,exist_ok=True)
    write_private(path,data)
    audit(db,actor,'migration',identifier,'SOURCE_STAGED',new={'file':filename,'sha256':file_hash,'rows':len(rows)})
    return {'id':identifier}


def require_migration(actor,roles={'AUTHOR','REVIEWER'}):
    if not any(g['role_code'] in roles and g['global_scope_flag'] for g in actor.grants): raise HTTPException(403,'A global content role is required for migration data.')

def batch_detail(db,actor,identifier):
    require_migration(actor)
    batch=repo.one(db,m.migration_batch,identifier)
    rows=[dict(r) for r in db.execute(sa.select(m.migration_row).where(m.migration_row.c.batch_id==identifier).order_by(m.migration_row.c.sheet_name,m.migration_row.c.row_number)).mappings()]
    for row in rows:row['profile']=json.loads(row['validation_errors'] or '{}');row['source']=json.loads(row['source_payload'])
    populations=Counter(r['validation_status'] for r in rows)
    return {'batch':batch,'rows':rows,'populations':dict(populations),'reconciled':len(rows)==batch['source_count']}

def validate_content(db,actor,identifier):
    require_migration(actor,{'REVIEWER'});batch=repo.one(db,m.migration_batch,identifier,lock=True)
    if actor.id==batch['uploaded_by']: raise HTTPException(403,'A different Compliance reviewer must validate migrated content.')
    count=0
    for row in db.execute(sa.select(m.migration_row).where(m.migration_row.c.batch_id==identifier,m.migration_row.c.validation_status=='CONTENT_VALIDATION_PENDING')).mappings():
        result=normalise(db,json.loads(row['source_payload']))
        if not result['errors'] and result['load_available']:
            result['validated_by']=actor.id;result['validated_at']=now().isoformat();count+=1
            db.execute(m.migration_row.update().where(m.migration_row.c.id==row['id']).values(validation_status='MIGRATION_READY',validation_errors=json.dumps(result)))
    audit(db,actor,'migration',identifier,'CONTENT_VALIDATED',new={'ready_rows':count})
    return {'validated':count}

def load_batch(db,actor,identifier):
    require_migration(actor,{'AUTHOR'});batch=repo.one(db,m.migration_batch,identifier,lock=True)
    rows=db.execute(sa.select(m.migration_row).where(m.migration_row.c.batch_id==identifier,m.migration_row.c.validation_status=='MIGRATION_READY').order_by(m.migration_row.c.row_number).with_for_update()).mappings().all()
    count=0
    for row in rows:
        result=normalise(db,json.loads(row['source_payload']))
        if result['errors']:raise HTTPException(409,f"Row {row['row_number']} changed eligibility: "+'; '.join(result['errors']))
        values=result['normalised']
        if result['entity_type']=='instrument':
            loaded=repo.create_instrument(db,actor,InstrumentInput(**values));entity='instrument'
        else:
            version=values.pop('version_id');loaded=repo.add_clause(db,actor,version,ClauseInput(**values),imported=True);entity='clause'
        db.execute(m.migration_row.update().where(m.migration_row.c.id==row['id']).values(validation_status='LOADED',target_entity_type=entity,target_entity_id=loaded['id'],loaded_at=now()))
        count+=1
    if not rows:raise HTTPException(409,'No migration-ready rows. Correct missing fields and obtain independent content validation first.')
    loaded=db.execute(sa.select(sa.func.count()).select_from(m.migration_row).where(m.migration_row.c.batch_id==identifier,m.migration_row.c.validation_status=='LOADED')).scalar_one()
    db.execute(m.migration_batch.update().where(m.migration_batch.c.id==identifier).values(loaded_count=loaded,pending_count=batch['source_count']-loaded,record_status='LOADED' if loaded==batch['source_count'] else 'PARTIALLY_LOADED',reconciliation_notes='Every loaded row retains its original file, sheet, row, payload, validation and target ID. All targets start Draft.'))
    audit(db,actor,'migration',identifier,'DRAFTS_LOADED',new={'loaded_now':count,'loaded_total':loaded})
    return {'loaded':count,'total':loaded}
